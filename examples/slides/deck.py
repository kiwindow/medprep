"""既存教材（HowToSetUpPC.pptx）の体裁に合わせたスライド組み立て道具。"""
import contextlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---- 既存デッキから採った色 -------------------------------------------------
INK   = RGBColor(0x13, 0x23, 0x3F)   # 本文（濃紺）
BLUE  = RGBColor(0x15, 0x65, 0xC0)   # 見出し・強調
CYAN  = RGBColor(0x1E, 0x90, 0xD8)   # バッジ
SKY   = RGBColor(0x00, 0xB0, 0xF0)   # 濃地の上の見出し
NAVY  = RGBColor(0x12, 0x33, 0x5C)   # 濃いカード
NAVY2 = RGBColor(0x0B, 0x25, 0x45)   # 表の見出し行
LIGHT = RGBColor(0xE8, 0xF1, 0xFB)   # 明るいカード
PALE  = RGBColor(0xF2, 0xF7, 0xFD)   # ごく明るいカード
MINT  = RGBColor(0xDF, 0xF3, 0xEE)   # 強調行
GREY  = RGBColor(0x5A, 0x6B, 0x84)   # 注記
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALEB = RGBColor(0xCF, 0xE0, 0xF5)   # 濃地の上の本文
RED   = RGBColor(0xE0, 0x31, 0x31)
ORANGE = RGBColor(0xE8, 0x59, 0x0C)
GREEN = RGBColor(0x1B, 0xAF, 0x7A)
CODEBG = RGBColor(0x0E, 0x21, 0x40)
CODEFG = RGBColor(0x6B, 0xD5, 0xFF)

FONT = "Yu Gothic"
MONO = "Consolas"

# ---- 文字の大きさ（既定より 2〜4pt 上げ、すべて Bold）-----------------------
T_TITLE = 26      # 大見出し（スライドタイトps）
T_EYE = 14        # eyebrow ラベル
T_CARD = 18       # 中見出し（カード見出し）
T_BODY = 15       # 本文
T_TABLE = 14      # 表のセル
T_NOTE = 13       # 注記


def _accent(color):
    """**…** の強調に使う色。すべて Bold なので、太さでは強調を表せない。"""
    if color == INK:
        return BLUE
    if color in (PALEB, RGBColor(0xC1, 0xE0, 0xF7)):
        return SKY
    if color == GREY:
        return INK
    return color


def _emph(text, color, size, mono):
    """`**…**` を色の変化に変える（PowerPoint に markdown は無い）。"""
    if "**" not in text:
        return [(text, color, size, mono)]
    out, acc = [], _accent(color)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            out.append((chunk, acc if i % 2 else color, size, mono))
    return out or [("", color, size, mono)]



def _width(text):
    """見た目の幅を「全角何文字ぶんか」で数える（半角は 0.55 文字ぶん）。"""
    n = 0.0
    for ch in text.replace("**", ""):
        n += 1.0 if ord(ch) > 0x2E80 else 0.55
    return n


def _rows(lines, cpl):
    """折り返しを含めた行数。"""
    return sum(max(1, -(-int(_width(t) * 100) // int(cpl * 100))) for t in lines)


def auto_h(lines, w, size, *, pad=0.32, head=0.0, line=1.36, extra=0.30):
    """本文の量から箱の高さを出す。★目分量で置かない。★"""
    if isinstance(lines, str):
        lines = [lines]
    # ★見積りは必ず余裕を取る。★ フォントが違えば折り返しは増える
    #   （Yu Gothic と代替フォントで 1 行ぶんずれる）。
    cpl = max(4.0, (w - pad * 2) * 72.0 / size * 0.92)
    return head + pad + _rows(lines, cpl) * size * line / 72.0 + extra


def code_h(lines, w, size, *, pad=0.26, extra=0.40):
    """等幅（ASCII）の箱の高さ。"""
    cpl = max(8.0, (w - pad * 2) * 72.0 / (size * 0.60))
    def cols(t):                      # 等幅では全角が 2 桁ぶん
        return sum(2 if ord(c) > 0x2E80 else 1 for c in t)
    n = sum(max(1, -(-cols(t) // int(cpl))) for t in lines)
    return pad + n * size * 1.35 / 72.0 + extra


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.3333)
        self.prs.slide_height = Inches(7.5)
        self.n = 0

    # ---------------------------------------------------------------- 骨
    def slide(self, eyebrow=None, title=None, number=True):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        if eyebrow:
            self.text(s, eyebrow, 0.55, 0.42, 11.0, 0.34, T_EYE, BLUE)
        if title:
            self.text(s, title, 0.55, 0.74, 12.23, 0.85, T_TITLE, INK)
        if number:
            self.text(s, str(self.n), 12.13, 7.08, 0.65, 0.30, 10, GREY,
                      align=PP_ALIGN.RIGHT)
        return s

    # ---------------------------------------------------------------- 部品
    def box(self, s, x, t, w, h, fill, radius=0.06):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(t), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        sh.shadow.inherit = False
        with contextlib.suppress(IndexError, KeyError):
            sh.adjustments[0] = radius
        sh.text_frame.text = ""
        return sh

    def text(self, s, runs, x, t, w, h, size, color, *, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, mono=False, line=1.25):
        tb = s.shapes.add_textbox(Inches(x), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if isinstance(runs, str):
            runs = [runs]
        for i, item in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            parts = item if isinstance(item, list) else _emph(item, color, size, mono)
            for part in parts:
                txt = part[0]
                col = part[1] if len(part) > 1 and part[1] else color
                sz = part[2] if len(part) > 2 and part[2] else size
                mn = part[3] if len(part) > 3 else mono
                r = p.add_run(); r.text = txt
                r.font.bold = True
                r.font.size = Pt(sz)
                r.font.color.rgb = col
                r.font.name = MONO if mn else FONT
                self._ea(r, MONO if mn else FONT)
        return tb

    @staticmethod
    def _ea(run, name):
        """日本語（East Asian）のフォントも明示する。"""
        rPr = run._r.get_or_add_rPr()
        from pptx.oxml.ns import qn
        for tag in ("a:latin", "a:ea", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {}); rPr.append(e)
            e.set("typeface", name)

    def badge(self, s, label, x, t, size=0.7, fill=BLUE, color=WHITE, fs=20):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(t),
                                Inches(size), Inches(size))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background(); sh.shadow.inherit = False
        tf = sh.text_frame; tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.bold = True; r.font.size = Pt(fs); r.font.color.rgb = color
        r.font.name = FONT; self._ea(r, FONT)
        return sh

    def card(self, s, x, t, w, h, head, body, *, fill=LIGHT, head_color=BLUE,
             body_color=INK, badge=None, badge_fill=BLUE, badge_color=WHITE,
             hs=T_CARD, bs=T_BODY):
        """h に None を渡すと、本文の量から高さを決める（推奨）。"""
        # 本文は t+0.95 から始まる。下の余白は 0.22。
        need = auto_h(body, w if badge is None else w - 0.82, bs,
                      head=0.95, pad=0.0, extra=0.22)
        h = need if h is None else max(h, need)
        self.box(s, x, t, w, h, fill)
        x = x + 0.32
        if badge is not None:
            self.badge(s, badge, x, t + 0.30, 0.62, badge_fill, badge_color, 18)
            x += 0.82
        self.text(s, head, x, t + 0.36, w - (x - x) - 0.30, 0.40, hs, head_color)
        self.text(s, body, x + 0.32, t + 0.95, w - 0.64, h - 1.15, bs, body_color)
        return t + h

    def callout(self, s, text, x, t, w, h=0.82, fill=NAVY, color=PALEB, size=T_BODY):
        flat = [t2[0] for row in text for t2 in row] if isinstance(text[0], list) else text
        need = auto_h(["".join(flat)] if isinstance(flat, list) else [flat],
                      w, size, pad=0.30, extra=0.30)
        h = max(h, need)
        self.box(s, x, t, w, h, fill)
        self.text(s, text, x + 0.30, t + 0.16, w - 0.60, h - 0.32, size, color,
                  anchor=MSO_ANCHOR.MIDDLE)
        return t + h

    def code(self, s, lines, x, t, w, h=None, size=13):
        """h に None を渡すと行数から高さを決める（推奨）。"""
        need = code_h(lines, w, size)
        h = need if h is None else max(h, need)
        self.box(s, x, t, w, h, CODEBG, radius=0.03)
        self.text(s, lines, x + 0.26, t + 0.20, w - 0.52, h - 0.40, size, CODEFG,
                  mono=True, line=1.35)
        return t + h

    def table(self, s, header, rows, x0, t, widths, *, rh=0.50, hh=0.56,
              colors=None, size=T_TABLE, hsize=T_TABLE):
        x = x0
        for i, htxt in enumerate(header):
            self.box(s, x, t, widths[i], hh, NAVY2, radius=0.02)
            self.text(s, htxt, x + 0.14, t, widths[i] - 0.28, hh, hsize, WHITE,
                      anchor=MSO_ANCHOR.MIDDLE)
            x += widths[i]
        y = t + hh
        for j, row in enumerate(rows):
            bg = PALE if j % 2 == 0 else WHITE
            if colors and colors.get(j):
                bg = colors[j]
            x = x0
            for i, cell in enumerate(row):
                self.box(s, x, y, widths[i], rh, bg, radius=0.02)
                col = INK if i else BLUE
                self.text(s, cell, x + 0.14, y, widths[i] - 0.28, rh, size, col,
                          anchor=MSO_ANCHOR.MIDDLE)
                x += widths[i]
            y += rh
        return y

    def note(self, s, text, x, t, w, h=0.34, color=GREY, size=T_NOTE):
        self.text(s, text, x, t, w, h, size, color)

    def save(self, path):
        self.prs.save(path)
        return path


def cover(deck, eyebrow, title_lines, subtitle, chips, footer):
    """表紙。既存デッキの濃紺＋幾何の形を踏襲する。"""
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                            deck.prs.slide_width, deck.prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0x05, 0x20, 0x3A)
    bg.line.fill.background(); bg.shadow.inherit = False
    for x, t, w, col in ((9.60, -1.60, 6.20, NAVY),
                         (10.70, -0.50, 4.00, RGBColor(0x12, 0x3A, 0x6A)),
                         (11.50, 0.40, 2.20, RGBColor(0x22, 0xAE, 0xFF))):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(t),
                               Inches(w), Inches(w))
        c.fill.solid(); c.fill.fore_color.rgb = col
        c.line.fill.background(); c.shadow.inherit = False
    deck.text(s, "mp", 11.50 + 0.75, 0.40 + 0.72, 1.2, 0.8, 34, RGBColor(0x05, 0x20, 0x3A),
              mono=True)
    deck.text(s, eyebrow, 0.55, 1.35, 9.0, 0.40, 16, CYAN)
    deck.text(s, title_lines, 0.55, 1.90, 9.6, 1.90, 44, WHITE, line=1.18)
    deck.text(s, subtitle, 0.55, 3.95, 10.0, 0.55, 22, PALEB)
    x = 0.55
    for chip in chips:
        w = 0.42 + 0.145 * len(chip)
        deck.box(s, x, 4.90, w, 0.50, RGBColor(0x13, 0x35, 0x60))
        deck.text(s, chip, x, 4.90, w, 0.50, 14, RGBColor(0xC1, 0xE0, 0xF7),
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w + 0.20
    deck.text(s, footer, 0.55, 5.95, 11.0, 0.40, 16, RGBColor(0x8C, 0xA4, 0xCC))
    return s
