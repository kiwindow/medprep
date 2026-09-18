"""論文に載る Table 1 / Table 2。**数値と書式が指示どおりであること**を確かめる。

ここで守りたい性質。

  1. Table 1 に p 値は出ない（背景を述べる表で検定はしない）
  2. 連続変数は 平均 ± SD [最小, 最大]、離散変数は n (%)
  3. 検定手法は表の中ではなく脚注に出る
  4. 群分けの指定が無ければ Table 2 は作らない
  5. 日本語版と英語版は**同じ数値**である（訳し直しで数字がずれない）
"""
import re

import numpy as np
import pandas as pd
import pytest

from medprep.schema import Schema
from medprep.tables import GTSummary, gt_tables

rng = np.random.default_rng(20260919)


def frame(n=120):
    return pd.DataFrame({
        "仮名ID": [f"P{i:04d}" for i in range(n)],
        "施設": rng.choice(["A院", "B院"], n),
        "性別": rng.choice(["男", "女"], n),
        "年齢": rng.normal(68, 12, n).round(0),
        "アルブミン(Alb)": rng.normal(3.6, 0.45, n).round(1),
    })


def built(df=None, group=None):
    df = frame() if df is None else df
    sch = Schema.infer(df, id_col="仮名ID", group="施設")
    return gt_tables(df, sch, group=group)


def test_table1_has_no_p_value():
    """★背景を述べる表で検定はしない。★"""
    gs = built()
    assert gs.table2 is None                       # 群分けが無ければ 1 枚だけ
    cols = list(gs.table1.frame_ja.columns)
    assert len(cols) == 2
    assert not any("p" in str(c).lower() for c in cols)


def test_continuous_is_mean_sd_min_max():
    """連続変数は **平均値 ± 標準偏差 [最小値, 最大値]**。"""
    df = frame()
    gs = built(df)
    row = gs.table1.frame_ja[gs.table1.frame_ja["特性"].str.startswith("年齢")]
    cell = row.iloc[0, 1]
    m = re.fullmatch(r"([\d,.-]+) ± ([\d,.-]+) \[([\d,.-]+), ([\d,.-]+)\]", cell)
    assert m, cell
    v = pd.to_numeric(df["年齢"])
    assert float(m.group(1).replace(",", "")) == pytest.approx(v.mean(), abs=0.6)
    assert float(m.group(3).replace(",", "")) == pytest.approx(v.min(), abs=0.6)
    assert float(m.group(4).replace(",", "")) == pytest.approx(v.max(), abs=0.6)


def test_categorical_is_n_percent():
    """離散変数は **n (%)**。"""
    df = frame()
    gs = built(df)
    row = gs.table1.frame_ja[gs.table1.frame_ja["特性"].str.startswith("性別")]
    assert len(row) == 1
    cell = row.iloc[0, 1]
    n_male = int((df["性別"] == "男").sum())
    assert cell.split(" / ")[0] == f"{n_male} ({100 * n_male / len(df):.1f}%)"


def test_binary_puts_the_positive_level_first():
    """`性別 — 女 / 男` では主語が分からない。**1 側を先に置く。**"""
    gs = built()
    labels = list(gs.table1.frame_ja.iloc[:, 0])
    assert any(x.startswith("性別 — 男 / 女") for x in labels), labels


def test_table2_has_p_on_the_right_and_tests_in_the_footnote():
    """★p 値は右端の 1 列だけ。検定手法は脚注。全体の列は入れない。★"""
    gs = built(group="施設")
    assert gs.table2 is not None
    cols = list(gs.table2.frame_ja.columns)
    assert cols[0] == "特性" and cols[-1].startswith("p値")
    # ★Table 2 に「全体」の列は入れない（全体は Table 1 の役目）★
    assert not any("全体" in c for c in cols)
    assert len(cols) == 2 + gs.table2.frame_ja.shape[1] - 2
    assert [c for c in cols if c.startswith("A院 (N = ")]
    assert [c for c in cols if c.startswith("B院 (N = ")]
    # 表の中には検定名が出ない
    body = gs.table2.frame_ja.astype(str).to_numpy().ravel()
    assert not any("検定" in x or "t 検定" in x for x in body)
    # 脚注には出る
    assert "検定" in gs.table2.footnote_ja or "分散分析" in gs.table2.footnote_ja


def test_japanese_and_english_hold_the_same_numbers():
    """**訳し直しで数字がずれない。**同じ数値から 2 版を作っている。"""
    gs = built(group="施設")
    for t in gs.tables():
        ja, en = t.frame_ja, t.frame_en
        assert ja.shape == en.shape
        # 見出し列以外は完全に同じ文字列
        assert (ja.iloc[:, 1:].to_numpy() == en.iloc[:, 1:].to_numpy()).all()
    assert "Albumin" in " ".join(gs.table1.frame_en.iloc[:, 0])
    assert "Sex — Male / Female" in list(gs.table1.frame_en.iloc[:, 0])


def test_internal_flag_columns_are_not_shown():
    """掃除の過程で付けた内部フラグは、人が読む表に出さない。"""
    df = frame()
    df["Alb__censored_low"] = 0
    df["欠損あり_年齢"] = 0
    gs = built(df)
    labels = " ".join(gs.table1.frame_ja.iloc[:, 0])
    assert "censored" not in labels and "欠損あり" not in labels


def test_excel_has_japanese_sheet1_and_english_sheet2(tmp_path):
    """★sheet1 = 日本語版、sheet2 = 英語版。★"""
    from openpyxl import load_workbook
    gs = built(group="施設")
    path = gs.to_excel(tmp_path / "t.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["日本語", "English"]
    assert "表1." in str(wb["日本語"]["A1"].value)
    assert "Table 1." in str(wb["English"]["A1"].value)


def test_word_separates_japanese_and_english_pages(tmp_path):
    """Word は日本語版と英語版をページを分けて載せる。表は中詰め。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    gs = built(group="施設")
    doc = Document(gs.to_docx(tmp_path / "t.docx"))
    assert len(doc.tables) == 4                       # 日 2 枚 + 英 2 枚
    assert len(doc.sections) == 4                     # それぞれ別ページ
    assert all(t.alignment == WD_TABLE_ALIGNMENT.CENTER for t in doc.tables)
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("表1.") for t in texts)
    assert any(t.startswith("Table 1.") for t in texts)


def test_html_is_centred_and_pasteable():
    """Word に貼れる HTML。配置は中詰め。"""
    gs = built()
    h = gs.to_html("ja")
    assert "<table" in h and "text-align:center" in h
    assert h.count("<tr>") == len(gs.table1.frame_ja) + 1      # 見出し行ぶん


def test_group_not_in_data_is_reported_not_crashed():
    gs = gt_tables(frame(), group="存在しない列")
    assert isinstance(gs, GTSummary) and gs.table2 is None
    assert any("存在しない列" in n for n in gs.notes)
