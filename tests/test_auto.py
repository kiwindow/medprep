"""層1（`mp.autoprep()`）。

確かめたいのは「1 行で走ること」ではない。次の 3 つである。

  1. **層1に固有のロジックが無い**（同じことを層2で 1 段ずつ再現できる）
  2. **判断がすべて外に出る**（`schema.yaml` と `rep.warnings`。空でないのがふつう）
  3. **リークしない**（train でしか fit していない）
"""
import os

import matplotlib
import numpy as np
import pandas as pd
import pytest

matplotlib.use("Agg")

import medprep as mp  # noqa: E402
from medprep import paths  # noqa: E402

rng = np.random.default_rng(20260918)


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "OUTPUT_DIR", str(tmp_path / "lab_output"))
    monkeypatch.setattr(paths, "WORK_DIR", str(tmp_path / "lab_work"))
    monkeypatch.setattr(paths, "IN_COLAB", False)
    monkeypatch.setattr(paths, "_LAST", {})
    return tmp_path


def frame(n=200):
    df = pd.DataFrame({
        "仮名ID": [f"P{i:04d}" for i in range(n)],
        "施設": rng.choice(["A院", "B院"], n),
        "性別": rng.choice(["男", "女"], n),
        "年齢": rng.normal(68, 12, n).round(0),
        "アルブミン(Alb)": rng.normal(3.6, 0.45, n).round(1),
        "末梢血｜血色素量(Hb)": rng.normal(10.8, 1.2, n).round(1),
        "無機リン(P)": rng.normal(5.2, 1.1, n).round(1),
        "転帰": (rng.random(n) < 0.3).astype(int),
    })
    df.loc[df.index[:15], "アルブミン(Alb)"] = np.nan
    return df


def run(df=None, **kw):
    kw.setdefault("verbose", False)
    return mp.autoprep(df if df is not None else frame(),
                       id_col="仮名ID", group="施設", **kw)


# ---------------------------------------------------------------- 骨格
def test_one_line_gets_a_model_ready_matrix():
    rep = run(outcome="転帰", task="classification")
    assert rep.X_train is not None and rep.X_test is not None
    assert list(rep.X_train.columns) == list(rep.X_test.columns)
    assert len(rep.X_train) + len(rep.X_test) == 200
    assert rep.pipeline is rep.prepared.preprocessor


def test_every_step_is_recorded_with_its_outcome():
    """★飛ばした段は黙って消さない。★"""
    rep = run(outcome="転帰", task="classification")
    names = [n for n, _ok, _d in rep.steps]
    for expected in ("列の役割を推定する", "品質を監査する", "辞書で掃除する",
                     "train / test に分ける", "リークを検査する"):
        assert expected in names


def test_without_an_outcome_it_stops_before_splitting():
    rep = run()
    assert rep.split is None and rep.prepared is None
    assert rep.schema is not None and rep.audit is not None


def test_quicklook_does_not_split():
    rep = mp.quicklook(frame(), id_col="仮名ID", group="施設", verbose=False)
    assert rep.split is None and rep.table1 is not None


def test_it_reads_a_file_and_says_what_it_guessed(tmp_path):
    p = tmp_path / "c.csv"
    frame().to_csv(p, index=False, encoding="cp932")
    rep = run(str(p))
    assert len(rep.df_raw) == 200
    assert any("cp932" in d for _n, _ok, d in rep.steps)


# ---------------------------------------------------------------- 判断の明示化
def test_the_warnings_are_not_empty_for_dirty_data():
    """★『確認事項なし』で返ってくるほうを疑うこと。★"""
    df = frame()
    df.loc[df.index[1], "仮名ID"] = df.loc[df.index[0], "仮名ID"]     # 同じ ID が 2 行
    rep = run(df, outcome="転帰", task="classification")
    assert rep.warnings
    assert any("重複" in w or "識別子" in w for w in rep.warnings)


def test_every_fatal_finding_reaches_the_warnings():
    df = frame()
    df.loc[df.index[1], "仮名ID"] = df.loc[df.index[0], "仮名ID"]
    rep = run(df)
    assert len(rep.warnings) >= len(rep.audit.errors)


def test_the_schema_can_be_written_out_and_read_back(tmp_path):
    rep = run(outcome="転帰", task="classification")
    p = tmp_path / "schema.yaml"
    rep.schema.to_yaml(p)
    back = mp.Schema.from_yaml(p)
    assert set(back.columns) == set(rep.schema.columns)


def test_dropping_missing_outcome_is_reported_not_silent():
    """★目的変数は補完してはならない。除いた数を必ず言う。★"""
    df = frame()
    df.loc[df.index[:20], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert len(rep.split.train) + len(rep.split.test) == 180
    assert any("目的変数" in w and "20" in w for w in rep.warnings)


# ---------------------------------------------------------------- リーク
def test_the_preprocessor_was_fitted_on_train_only():
    rep = run(outcome="転帰", task="classification")
    lc = mp.leak_check(rep.pipeline, rep.split.train, rep.split.test)
    assert (lc["結果"] == "OK").all()


def test_the_identifier_never_becomes_a_feature():
    rep = run(outcome="転帰", task="classification")
    assert "仮名ID" not in rep.X_train.columns
    assert "転帰" not in rep.X_train.columns


# ---------------------------------------------------------------- 飛ばし方
def test_a_failing_optional_step_is_skipped_with_its_reason(monkeypatch):
    """付随する段が落ちても止まらない。ただし理由を残す。"""
    from medprep import auto
    monkeypatch.setattr(auto, "target_achievement",
                        lambda *a, **k: (_ for _ in ()).throw(ValueError("わざと")))
    rep = run()
    skipped = [(n, d) for n, ok, d in rep.steps if not ok]
    assert any("管理目標" in n and "わざと" in d for n, d in skipped)


def test_a_failing_backbone_step_stops_everything(monkeypatch):
    """★背骨で失敗したら止まる。前処理だけ抜けた結果を返してはならない。★"""
    from medprep import auto
    monkeypatch.setattr(auto, "prepare",
                        lambda *a, **k: (_ for _ in ()).throw(ValueError("止まる")))
    with pytest.raises(ValueError, match="止まる"):
        run(outcome="転帰", task="classification")


# ---------------------------------------------------------------- 保存
def test_save_writes_the_run_folder(isolated):
    rep = run(outcome="転帰", task="classification", save=True)
    p = rep.run
    assert p.runnumber == 1
    for rel in (("model", "schema.yaml"), ("model", "pipeline.pkl"),
                ("model", "medprep_version.txt"), ("table", "run_info.json"),
                ("table", "prep_tables.xlsx"), ("report", "prep_report.html")):
        assert os.path.exists(p.file(*rel)), rel
    assert len(os.listdir(p.figure)) >= 5


def test_the_run_folder_is_gitignored(isolated):
    """★出力には個票由来の値が入りうる。git に載せない。★"""
    run(save=True)
    assert (isolated / "lab_output" / ".gitignore").exists()


def test_two_runs_do_not_collide(isolated):
    a = run(save=True)
    b = run(save=True)
    assert a.run.runnumber != b.run.runnumber


def test_the_log_records_the_run(isolated):
    run(outcome="転帰", task="classification", save=True)
    hist = isolated / "lab_work" / "log" / "history.csv"
    assert hist.exists()
    row = pd.read_csv(hist).iloc[0]
    assert row["手法"] == "Preprocessing" and row["状態"] == "完了"


def test_nothing_is_written_without_save(isolated):
    run(outcome="転帰", task="classification")
    assert not (isolated / "lab_output").exists()


# ---------------------------------------------------------------- レポート
def test_case_level_values_are_hidden_by_default():
    df = frame()
    df.loc[df.index[1], "仮名ID"] = df.loc[df.index[0], "仮名ID"]
    rep = run(df)
    assert df.loc[df.index[0], "仮名ID"] not in rep.html.to_html()
    assert df.loc[df.index[0], "仮名ID"] in rep.to_html(show_values=True).to_html()


def test_the_report_leads_with_the_warnings():
    df = frame()
    df.loc[df.index[1], "仮名ID"] = df.loc[df.index[0], "仮名ID"]
    rep = run(df)
    html = rep.html.to_html()
    assert html.index("まず読むこと") < html.index("データ品質監査")


def test_report_text_lists_the_steps_and_the_warnings():
    rep = run(outcome="転帰", task="classification")
    text = rep.report()
    assert "段:" in text and "人の確認が要る事項" in text


def test_to_excel_writes_one_sheet_per_table(tmp_path):
    rep = run()
    p = rep.to_excel(tmp_path / "t.xlsx")
    assert len(pd.read_excel(p, sheet_name=None)) >= 1


def test_the_survival_frame_carries_the_covariates():
    """★共変量の無い SurvivalFrame は Cox に掛けられない。★

    ID と (duration, event) だけを返すと、受け取った側が元データと
    結合し直すはめになり、そこで症例を取り違える。
    """
    df = frame()
    base = pd.Timestamp("2015-01-01")
    df["開始日"] = base + pd.to_timedelta(rng.integers(0, 500, len(df)), "D")
    df["発生日"] = df["開始日"] + pd.to_timedelta(rng.integers(30, 900, len(df)), "D")
    df["打切日"] = ""
    df.loc[df.index[::2], ["発生日"]] = ""
    df.loc[df.index[::2], "打切日"] = (
        df.loc[df.index[::2], "開始日"] + pd.Timedelta(days=700))
    rep = run(df, survival_dates=("開始日", "発生日", "打切日"))
    cols = set(rep.survival.data.columns)
    assert {"duration", "event"} <= cols
    assert {"年齢", "施設"} <= cols


def _survival_frame(n=200):
    df = frame(n)
    base = pd.Timestamp("2015-01-01")
    df["開始日"] = base + pd.to_timedelta(rng.integers(0, 500, len(df)), "D")
    df["発生日"] = df["開始日"] + pd.to_timedelta(rng.integers(30, 900, len(df)), "D")
    df["打切日"] = ""
    df.loc[df.index[::2], "発生日"] = ""
    df.loc[df.index[::2], "打切日"] = (
        df.loc[df.index[::2], "開始日"] + pd.Timedelta(days=700))
    return df


def test_duration_and_event_are_written_into_the_data_not_only_into_the_report():
    """★3 列の日付を指定したのに、出来上がったデータに観察期間が無いのでは
    指定した意味がない。★

    `rep.survival.data` の中にしか無いと、Excel を開いた人には見えない。
    掃除済み・解析用データの両方に `duration` / `event` / `_start` / `_end` を残す。
    """
    rep = run(_survival_frame(), survival_dates=("開始日", "発生日", "打切日"))
    for c in ("duration", "event", "_start", "_end"):
        assert c in rep.df_clean.columns, c
        assert c in rep.df_use.columns, c
    assert any("観察期間とイベント" in n for n, _ok, _d in rep.steps)

    # 値が SurvivalFrame と一致する（結合を取り違えていない）
    sf = rep.survival.data
    got = rep.df_clean.loc[sf.index, "duration"]
    assert (got - sf["duration"]).abs().max() < 1e-9


def test_survival_columns_do_not_become_features():
    """★目的変数そのものが説明変数に紛れ込んではならない。★

    `duration` と `event` を残すのは人が読むためであって、モデルに入れるためではない。
    """
    rep = run(_survival_frame(), survival_dates=("開始日", "発生日", "打切日"))
    feats = rep.schema.features()
    assert "duration" not in feats and "event" not in feats
    assert rep.schema.columns["duration"].role == "time"
    assert rep.schema.columns["event"].role == "event"


def test_cases_that_could_not_be_converted_are_nan_not_dropped():
    """★行は削除しない。★ 生存時間にできなかった症例は NaN にする。"""
    df = _survival_frame()
    df.loc[df.index[:10], "開始日"] = ""          # 開始日なし → 変換できない
    rep = run(df, survival_dates=("開始日", "発生日", "打切日"))
    assert len(rep.df_clean) == len(df)                     # 行は減らない
    assert rep.df_clean["duration"].isna().sum() >= 10


# ---------------------------------------------------------------- 減らしたもの
def test_it_records_what_it_removed():
    """★何を捨てたかを言わない自動化は、信用してはならない。★"""
    df = frame()
    df.loc[df.index[:12], "転帰"] = np.nan          # 目的変数の欠測 → 行が減る
    df["連番"] = range(len(df))                     # 使われない列
    rep = run(df, outcome="転帰", task="classification")
    r = rep.removed
    assert set(r.columns) == {"種類", "対象", "件数", "処置", "理由", "段"}
    assert (r["理由"].astype(str).str.len() > 0).all()      # 理由の無い行を作らない

    # 列：ID は解析から外れ、その理由が書いてある
    col = r[r["種類"] == "列"]
    assert "仮名ID" in set(col["対象"])
    assert "id_col" in col.loc[col["対象"] == "仮名ID", "理由"].iloc[0]

    # 行：目的変数の欠測に**印を付けた**数が一致する（削除はしていない）
    row = r[(r["種類"] == "行")]
    assert int(row["件数"].sum()) == 12
    assert rep.n_excluded == 12
    assert len(rep.df_clean) == len(df)


def test_removed_is_empty_when_nothing_was_removed():
    df = pd.DataFrame({"年齢": [60, 70, 80], "Alb": [3.1, 3.5, 4.0]})
    rep = mp.autoprep(df, verbose=False)
    assert rep.removed is not None and len(rep.removed) == 0


def test_cleaned_values_are_counted_as_removed():
    """辞書で NaN にした値も「減らしたもの」である（行は消えないが、値は消える）。"""
    df = frame()
    df["インタクトPTH(iPTH)"] = 999.0                # 欠損コード
    rep = run(df)
    val = rep.removed[rep.removed["種類"] == "値"]
    assert len(val) and "999" in " ".join(val["理由"])


def test_the_report_text_lists_what_was_removed():
    df = frame()
    df.loc[df.index[:5], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert "減らしたもの" in rep.report()


def test_the_html_report_has_a_removed_section():
    df = frame()
    df.loc[df.index[:5], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert "減らしたもの" in rep.html.to_html()


# ---------------------------------------------------------------- データの書き出し
def test_save_writes_the_cleaned_and_prepared_data(isolated):
    rep = run(outcome="転帰", task="classification", save=True)
    d = os.path.join(rep.run.run, "data")
    for name in ("掃除済みデータ.xlsx", "前処理済み_train.xlsx", "前処理済み_test.xlsx"):
        assert os.path.exists(os.path.join(d, name)), name
    assert os.path.exists(rep.run.file("table", "除外の記録.xlsx"))


def test_the_prepared_file_carries_the_outcome_column(isolated):
    rep = run(outcome="転帰", task="classification", save=True)
    tr = pd.read_excel(os.path.join(rep.run.run, "data", "前処理済み_train.xlsx"))
    assert "転帰" in tr.columns
    assert len(tr) == len(rep.X_train)
    # 先頭に「元の行」「仮名ID」、末尾に目的変数。その間が特徴量。
    assert list(tr.columns)[2:-1] == list(rep.X_train.columns)


def test_save_data_can_be_turned_off(isolated):
    """★症例レベルのデータである。要らないなら書かない。★"""
    rep = run(outcome="転帰", task="classification", save=True, save_data=False)
    assert not os.path.exists(os.path.join(rep.run.run, "data"))


# ---------------------------------------------------------------- 日付
def messy_dates(n=80):
    """和暦・全角・Excel シリアル値・時刻付きが 1 列に混ざったデータ。"""
    base = pd.Timestamp("2015-01-01")
    days = rng.integers(0, 2000, n)
    ts = [base + pd.Timedelta(days=int(d)) for d in days]
    styles = ["ymd", "wareki", "zenkaku", "serial", "time"] * (n // 5 + 1)
    out = []
    for t, st in zip(ts, styles[:n]):
        if st == "ymd":
            out.append(f"{t.year}/{t.month}/{t.day}")
        elif st == "wareki":
            out.append(f"H{t.year - 1988}.{t.month}.{t.day}" if t.year < 2019
                       else f"R{t.year - 2018}.{t.month}.{t.day}")
        elif st == "zenkaku":
            out.append(f"{t.year}/{t.month}/{t.day}".translate(
                str.maketrans("0123456789/", "０１２３４５６７８９／")))
        elif st == "serial":
            out.append(int((t - pd.Timestamp("1899-12-30")).days))
        else:
            out.append(f"{t.year}-{t.month}-{t.day} 09:40:38")
    df = frame(n)
    df["検査日"] = out
    return df


def test_messy_dates_are_parsed_into_real_dates():
    """★掃除は数値列にしか掛からない。日付を直さなければ生の文字列のまま残る。★"""
    rep = run(messy_dates(), date_col="検査日")
    col = rep.df_clean["検査日"]
    assert pd.api.types.is_datetime64_any_dtype(col)
    assert col.notna().all()
    assert rep.dates["検査日"].order == "ymd"


def test_the_parsing_is_reported_as_a_step():
    rep = run(messy_dates(), date_col="検査日")
    assert any(n == "日付を解釈する" and ok for n, ok, _d in rep.steps)


def test_an_ambiguous_date_column_is_left_alone():
    """★日と月の順序が決まらない列に手を触れない。★

    3/4 が 3月4日なのか 4月3日なのか分からないまま観察期間を計算してはならない。
    """
    df = frame(60)
    df["検査日"] = ["3/4/15", "5/6/15", "7/8/15", "2/1/15"] * 15
    rep = run(df, date_col="検査日")
    assert rep.df_clean["検査日"].dtype == object          # 書き換えていない
    assert any("順序が決まらない" in w for w in rep.warnings)


def test_a_column_that_is_not_a_date_is_left_alone():
    df = frame(60)
    df["メモ"] = ["特記なし"] * 60
    rep = run(df, date_col="メモ")
    assert rep.df_clean["メモ"].dtype == object


def test_parsed_dates_reach_the_saved_file(isolated):
    rep = run(messy_dates(), date_col="検査日", save=True)
    book = pd.read_excel(os.path.join(rep.run.run, "data", "掃除済みデータ.xlsx"),
                         sheet_name="データ")
    assert pd.api.types.is_datetime64_any_dtype(book["検査日"])


# ---------------------------------------------------------------- 保存の形
def test_the_cleaned_workbook_says_which_columns_were_dropped(isolated):
    """★落とすと決めた列が、保存したファイルから分かること。★"""
    rep = run(save=True)
    book = pd.read_excel(os.path.join(rep.run.run, "data", "掃除済みデータ.xlsx"),
                         sheet_name=None)
    assert set(book) == {"データ", "列の扱い", "減らしたもの"}
    treat = book["列の扱い"]
    assert "仮名ID" in set(treat["列名"])
    assert treat.loc[treat["列名"] == "仮名ID", "扱い"].iloc[0] == "drop"


def test_the_analysis_file_has_the_dropped_columns_removed(isolated):
    rep = run(save=True)
    d = os.path.join(rep.run.run, "data")
    full = pd.read_excel(os.path.join(d, "掃除済みデータ.xlsx"), sheet_name="データ")
    use = pd.read_excel(os.path.join(d, "解析用データ.xlsx"))
    assert "仮名ID" in full.columns          # 人が症例を辿れるように残す
    assert "仮名ID" not in use.columns       # 解析には使わない
    assert len(use) == len(full)             # 行は減らさない


# ---------------------------------------------------------------- 行は消さない
def test_rows_are_marked_not_deleted():
    """★行を削除すると、元のデータと症例ごとに axis=1 で結合し直せなくなる。★"""
    df = frame(200)
    df.loc[df.index[:15], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert len(rep.df_clean) == len(df)                 # 行数が変わらない
    assert list(rep.df_clean.index) == list(df.index)   # 並びも変わらない
    assert rep.n_excluded == 15
    assert rep.df_clean["除外推奨"].sum() == 15


def test_the_mark_carries_its_reason():
    df = frame(120)
    df.loc[df.index[:8], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    marked = rep.df_clean.loc[rep.df_clean["除外推奨"] == 1, "除外推奨_理由"]
    assert len(marked) == 8
    assert marked.str.contains("目的変数").all()


def test_the_cleaned_frame_can_be_concatenated_with_the_original():
    """★これが行を消さない理由そのもの。★"""
    df = frame(150)
    df.loc[df.index[:10], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    both = pd.concat([df, rep.df_clean[["除外推奨", "除外推奨_理由"]]], axis=1)
    assert len(both) == len(df)
    assert both["除外推奨"].notna().all()


def test_the_mark_columns_are_not_features():
    """印は人が読むための列である。特徴量にしてはならない。"""
    df = frame(150)
    df.loc[df.index[:10], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert "除外推奨" not in rep.X_train.columns
    assert "除外推奨_理由" not in rep.X_train.columns


def test_the_model_matrix_still_excludes_the_marked_rows():
    """モデルには渡せない（目的変数が欠測なので）。そこは部分集合になる。"""
    df = frame(150)
    df.loc[df.index[:10], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    assert len(rep.X_train) + len(rep.X_test) == 140


def test_the_saved_files_keep_every_row(isolated):
    df = frame(150)
    df.loc[df.index[:10], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification", save=True)
    d = os.path.join(rep.run.run, "data")
    for name in ("掃除済みデータ.xlsx", "解析用データ.xlsx"):
        assert len(pd.read_excel(os.path.join(d, name), sheet_name=0)) == 150, name
    use = pd.read_excel(os.path.join(d, "解析用データ.xlsx"))
    assert {"除外推奨", "除外推奨_理由"} <= set(use.columns)


def test_the_prepared_file_can_be_traced_back_to_the_original_rows(isolated):
    """train / test は部分集合なので、元の行番号と ID を付けておく。"""
    rep = run(outcome="転帰", task="classification", save=True)
    tr = pd.read_excel(os.path.join(rep.run.run, "data", "前処理済み_train.xlsx"))
    assert list(tr.columns)[:2] == ["元の行", "仮名ID"]
    assert set(tr["元の行"]) <= set(rep.df_clean.index)


def test_removed_says_the_rows_were_only_marked():
    df = frame(120)
    df.loc[df.index[:6], "転帰"] = np.nan
    rep = run(df, outcome="転帰", task="classification")
    row = rep.removed[rep.removed["種類"] == "行"]
    assert (row["処置"].str.contains("削除しない")).all()
    assert "削除していない" in rep.report()


def test_dates_are_written_without_the_time_part(isolated):
    """★Excel に 00:00:00 を出さない。★ 検査日は年月日で読むものである。"""
    import openpyxl
    rep = run(messy_dates(), date_col="検査日", save=True)
    path = os.path.join(rep.run.run, "data", "掃除済みデータ.xlsx")
    ws = openpyxl.load_workbook(path)["データ"]
    header = [c.value for c in ws[1]]
    col = header.index("検査日") + 1
    cell = ws.cell(row=2, column=col)
    # ★Excel の表示を決めるのは書式である。★ 時刻を含む書式にしない。
    assert cell.number_format == "YYYY-MM-DD"
    assert "HH" not in cell.number_format and "SS" not in cell.number_format
    # 計算に使うほうは datetime のまま（日付の引き算に要る）
    assert pd.api.types.is_datetime64_any_dtype(rep.df_clean["検査日"])


def test_the_outputs_table_explains_each_file(isolated):
    """★3 つがどう違うのか、何をしたのかが表で分かること。★"""
    rep = run(outcome="転帰", task="classification", save=True)
    o = rep.outputs
    assert set(o.columns) == {"ファイル", "行", "列", "施した処置", "使いどころ"}
    names = " ".join(o["ファイル"])
    assert "掃除済み" in names and "解析用" in names and "前処理済み" in names
    # 段が進むほど処置が増える
    lens = o["施した処置"].str.len().tolist()
    assert lens[0] < lens[1] < lens[2]
    assert os.path.exists(rep.run.file("table", "書き出したデータの説明.xlsx"))


def test_the_report_explains_the_three_files(isolated):
    rep = run(outcome="転帰", task="classification", save=True)
    html = rep.html.to_html()
    assert "書き出したデータ" in html
    assert "axis=1" in html


def test_outputs_is_empty_without_saving():
    rep = run(outcome="転帰", task="classification")
    assert rep.outputs is not None and len(rep.outputs) == 0


# ---------------------------------------------------------------- 書き出し書式
def test_small_numbers_keep_their_decimals_in_excel(tmp_path):
    """★小さい値が 0 に見えてはいけない。★

    CRP は 0.01〜13 の幅がある。既定の書式のままだと、表示環境によっては
    0.05 が「0」に見える。列ごとに、いちばん小さい非ゼロの値が読める桁数を当てる。
    """
    from openpyxl import load_workbook

    from medprep.auto import _decimals, _excel

    assert _decimals([0.01, 0.05, 2.96, 12.97]) == 2      # ★最低 2 桁★
    assert _decimals([65, 70, 80]) == 0
    assert _decimals([2.1, 3.6, 4.8]) == 2
    assert _decimals([0.0027, 1.41, 9.2]) == 3            # 0.00 に丸めない

    d = pd.DataFrame({"CRP": [0.01, 0.05, 2.96], "年齢": [65, 70, 80]})
    path = tmp_path / "t.xlsx"
    _excel(d, path)
    ws = load_workbook(path).active
    assert ws.cell(2, 1).number_format == "0.00"
    assert ws.cell(2, 2).number_format == "0"


def test_the_prepared_matrix_also_keeps_the_raw_values_on_a_second_sheet(tmp_path):
    """★標準化後の表は人が読んでも意味が取れない。★

    施設は `施設=B院` のダミーになり、数値は z 値になる。群間比較をしたり
    値を目で確かめたりできるよう、同じ症例・同じ列を**生の値**でも残す。
    """
    from openpyxl import load_workbook

    df = _survival_frame()
    run(df, outcome="転帰", task="classification",
        save=True, out_dir=str(tmp_path), method="T")
    book = tmp_path / "T" / "run1" / "data" / "前処理済み_train.xlsx"
    assert book.exists()
    assert load_workbook(book).sheetnames == ["標準化後", "生の値", "欠損値の位置"]

    raw = pd.read_excel(book, sheet_name="生の値")
    std = pd.read_excel(book, sheet_name="標準化後")
    assert len(raw) == len(std)
    assert "施設" in raw.columns                      # 名義尺度のまま
    assert set(raw["施設"].dropna()) <= {"A院", "B院"}
    assert not [c for c in raw.columns if c.startswith("施設=")]
    assert [c for c in std.columns if c.startswith("施設=")]   # こちらはダミー


def test_the_untouched_original_is_saved_too(tmp_path):
    """★掃除の前と後を突き合わせられないと、機械が直したのか分からなくなる。★"""
    df = _survival_frame()
    rep = run(df, outcome="転帰", task="classification",
              save=True, out_dir=str(tmp_path), method="T")
    raw = pd.read_excel(tmp_path / "T" / "run1" / "data" / "元データ.xlsx")
    assert len(raw) == len(df)
    assert list(raw.columns) == list(df.columns)
    assert any("元データ" in f for f in rep.saved)


def test_it_says_why_there_is_no_train_test_when_no_outcome_was_given():
    """★飛ばした段を黙って消さない。★

    目的変数が無ければ分割にも標準化にも意味がないので作らない。しかし黙って
    消すと「前処理済み_train.xlsx が無い」とだけ見えて、**理由が分からなくなる**。
    """
    rep = run(frame())                     # outcome を渡さない
    entry = [(n, d) for n, _ok, d in rep.steps if n == "train / test に分ける"]
    assert entry, "段そのものが記録に残っていない"
    assert "目的変数" in entry[0][1] and "前処理済み" in entry[0][1]
    assert any("前処理済みの行列" in n for n in rep.notes)
    assert rep.prepared is None


def test_the_outputs_table_lists_the_file_it_did_not_make(tmp_path):
    """作らなかったファイルも表に出す。**空欄は「無い」の説明にならない。**"""
    run(frame(), save=True, out_dir=str(tmp_path), method="T")
    import glob
    book = glob.glob(str(tmp_path / "T" / "run1" / "table" / "書き出したデータの説明.xlsx"))
    assert book
    t = pd.read_excel(book[0])
    names = " ".join(t["ファイル"].astype(str))
    assert "元データ.xlsx" in names
    assert "前処理済み_train.xlsx" in names
    row = t[t["ファイル"].astype(str).str.contains("前処理済み")].iloc[0]
    assert "作っていない" in str(row["施した処置"])


def test_the_cases_flagged_for_exclusion_are_listed_in_their_own_sheet(tmp_path):
    """★600 行の中から 22 行を探すのは人の仕事ではない。★

    印は掃除済みデータの 2 列にも入っているが、それとは別に一覧を出す。
    理由だけでなく**日付の列も添える** ―― 除外の理由はほとんど日付の矛盾なので、
    理由だけ見せられても人は確かめようがない。
    """
    from openpyxl import load_workbook

    import medprep as mp

    df = _survival_frame(300)
    df.loc[df.index[:8], "開始日"] = ""           # 開始日なし → 除外推奨
    rep = run(df, survival_dates=("開始日", "発生日", "打切日"),
              save=True, out_dir=str(tmp_path), method="T")
    assert rep.n_excluded >= 8
    assert len(rep.df_clean) == len(df)            # ★行は削除しない★

    book = tmp_path / "T" / "run1" / "table" / "除外の記録.xlsx"
    assert load_workbook(book).sheetnames == ["まとめ", "外すのが望ましい症例"]
    cases = pd.read_excel(book, sheet_name="外すのが望ましい症例")
    assert len(cases) == rep.n_excluded
    assert {"元の行", "仮名ID", "理由"} <= set(cases.columns)
    assert (cases["理由"].astype(str).str.len() > 0).all()
    assert "開始日" in cases.columns                # 理由を確かめる材料が添えてある

    # 同じものが API からも取れる
    assert len(mp.excluded_cases(rep)) == rep.n_excluded
